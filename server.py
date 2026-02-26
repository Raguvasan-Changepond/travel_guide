#!/usr/bin/env python3
"""
Web server for Travel Guide.

Runs both static frontend and JSON APIs:
- Authentication/session APIs
- Agenda Excel API
- JSON config APIs
"""

from http.cookies import SimpleCookie
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path
from datetime import datetime, timedelta, timezone
import hashlib
import hmac
import json
import openpyxl
import os
import re
import secrets
import threading
import time
import urllib.parse


BASE_DIR = Path(__file__).resolve().parent
CONFIG_DIR = BASE_DIR / "config"
PUBLIC_DIR = BASE_DIR / "public"
DATA_DIR = CONFIG_DIR / "data"
AGENDA_EXTENSIONS = {".xlsx", ".xls"}


def resolve_latest_agenda_path():
    """Return the most recently modified agenda Excel file from config/."""
    candidates = [
        path
        for path in CONFIG_DIR.iterdir()
        if path.is_file() and path.suffix.lower() in AGENDA_EXTENSIONS
    ]
    if not candidates:
        raise FileNotFoundError(f"No agenda Excel file found in: {CONFIG_DIR}")
    return max(candidates, key=lambda path: path.stat().st_mtime)


def detect_header_row(ws, search_rows=10):
    """Find the first likely header row by scanning for non-empty cells."""
    upper_bound = min(ws.max_row, search_rows)
    for row_idx in range(1, upper_bound + 1):
        non_empty = 0
        for col_idx in range(1, ws.max_column + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is not None and str(value).strip():
                non_empty += 1
        if non_empty >= 2:
            return row_idx
    return 1


def normalize_label(value):
    """Normalize sheet/header labels for robust matching."""
    if value is None:
        return ""
    return str(value).strip().lower()


def read_env_int(name, default):
    """Read integer env var with safe fallback."""
    try:
        return int(os.getenv(name, str(default)))
    except (TypeError, ValueError):
        return default


class DynamicExcelHandler(SimpleHTTPRequestHandler):
    """Serve static files and secure API endpoints."""

    SESSION_COOKIE_NAME = "travelguide_session"
    SESSION_TTL_MINUTES = read_env_int("SESSION_TTL_MINUTES", 480)
    LOGIN_WINDOW_SECONDS = read_env_int("LOGIN_WINDOW_SECONDS", 300)
    LOGIN_MAX_ATTEMPTS = read_env_int("LOGIN_MAX_ATTEMPTS", 8)
    COOKIE_SECURE = os.getenv("COOKIE_SECURE", "").lower() in {"1", "true", "yes", "on"}

    SESSIONS = {}
    FAILED_LOGINS = {}
    AUTH_LOCK = threading.Lock()

    CONFIG_ENDPOINTS = {
        "/api/settings": DATA_DIR / "settings.json",
        "/api/about": DATA_DIR / "about.json",
        "/api/places": DATA_DIR / "places.json",
    }

    def handle(self):
        """Ignore common browser disconnects to avoid noisy stack traces."""
        try:
            super().handle()
        except (BrokenPipeError, ConnectionResetError, ConnectionAbortedError):
            return

    def translate_path(self, path):
        """Serve files only from public directory."""
        parsed_path = urllib.parse.urlparse(path).path
        if parsed_path in {"", "/"}:
            parsed_path = "/index.html"

        requested = (PUBLIC_DIR / urllib.parse.unquote(parsed_path).lstrip("/")).resolve()
        public_root = PUBLIC_DIR.resolve()

        if not str(requested).startswith(str(public_root)):
            return str(public_root / "index.html")
        return str(requested)

    def end_headers(self):
        """Add baseline security headers."""
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("X-Frame-Options", "DENY")
        self.send_header("Referrer-Policy", "strict-origin-when-cross-origin")
        self.send_header("Permissions-Policy", "geolocation=(), microphone=(), camera=()")
        super().end_headers()

    def do_GET(self):
        """Handle GET routes."""
        route = urllib.parse.urlparse(self.path).path

        # Browsers probe /favicon.ico by default; map it to our configured favicon.
        if route == "/favicon.ico":
            self.path = "/assets/images/cp_logo_fav_iccon.jfif"
            try:
                return self.serve_static_request(include_body=True)
            except (BrokenPipeError, ConnectionResetError, ConnectionAbortedError):
                return

        if route == "/api/health":
            self.send_json(200, {"status": "ok"})
            return
        if route == "/api/session":
            self.handle_session_status()
            return
        if route in self.CONFIG_ENDPOINTS:
            if not self.require_auth():
                return
            self.handle_config_request(route)
            return
        if route == "/api/agenda":
            if not self.require_auth():
                return
            self.handle_agenda_request()
            return
        if route.startswith("/api/"):
            self.send_json(404, {"error": "Endpoint not found"})
            return

        try:
            return self.serve_static_request(include_body=True)
        except (BrokenPipeError, ConnectionResetError, ConnectionAbortedError):
            return

    def do_HEAD(self):
        """Handle HEAD routes for static files."""
        route = urllib.parse.urlparse(self.path).path
        if route.startswith("/api/"):
            self.send_json(404, {"error": "Endpoint not found"})
            return
        try:
            return self.serve_static_request(include_body=False)
        except (BrokenPipeError, ConnectionResetError, ConnectionAbortedError):
            return

    def do_POST(self):
        """Handle POST routes."""
        route = urllib.parse.urlparse(self.path).path

        if route == "/api/login":
            self.handle_login()
            return
        if route == "/api/logout":
            self.handle_logout()
            return
        if route.startswith("/api/"):
            self.send_json(404, {"error": "Endpoint not found"})
            return

    def serve_static_request(self, include_body=True):
        """Serve static files with byte-range support for media seeking."""
        route = urllib.parse.urlparse(self.path).path
        requested_path = Path(self.translate_path(route)).resolve()
        public_root = PUBLIC_DIR.resolve()

        if not str(requested_path).startswith(str(public_root)):
            self.send_error(403, "Forbidden")
            return

        if requested_path.is_dir():
            requested_path = (requested_path / "index.html").resolve()

        if not requested_path.exists() or not requested_path.is_file():
            self.send_error(404, "File not found")
            return

        file_size = requested_path.stat().st_size
        range_header = self.headers.get("Range")
        start = 0
        end = file_size - 1
        status_code = 200

        if range_header:
            match = re.match(r"bytes=(\d*)-(\d*)$", range_header.strip())
            if not match:
                self.send_error(416, "Invalid Range")
                return

            start_str, end_str = match.groups()
            if start_str == "" and end_str == "":
                self.send_error(416, "Invalid Range")
                return

            if start_str == "":
                suffix_length = int(end_str)
                if suffix_length <= 0:
                    self.send_error(416, "Invalid Range")
                    return
                start = max(file_size - suffix_length, 0)
            else:
                start = int(start_str)
                if start >= file_size:
                    self.send_error(416, "Range Not Satisfiable")
                    return

            if end_str != "":
                end = min(int(end_str), file_size - 1)
            else:
                end = file_size - 1

            if end < start:
                self.send_error(416, "Range Not Satisfiable")
                return

            status_code = 206

        content_length = end - start + 1
        content_type = self.guess_type(str(requested_path))

        self.send_response(status_code)
        self.send_header("Content-Type", content_type)
        self.send_header("Accept-Ranges", "bytes")
        self.send_header("Content-Length", str(content_length))
        if status_code == 206:
            self.send_header("Content-Range", f"bytes {start}-{end}/{file_size}")
        self.end_headers()

        if not include_body:
            return

        with requested_path.open("rb") as file_obj:
            file_obj.seek(start)
            remaining = content_length
            chunk_size = 64 * 1024
            while remaining > 0:
                chunk = file_obj.read(min(chunk_size, remaining))
                if not chunk:
                    break
                self.wfile.write(chunk)
                remaining -= len(chunk)

        self.send_error(405, "Method Not Allowed")

    def handle_login(self):
        """Authenticate and create secure session."""
        client_ip = self.get_client_ip()
        if self.is_rate_limited(client_ip):
            self.send_json(429, {"error": "Too many failed login attempts. Please try again shortly."})
            return

        try:
            payload = self.read_json_body()
        except ValueError as err:
            self.send_json(400, {"error": str(err)})
            return

        username = str(payload.get("username", "")).strip()
        password = str(payload.get("password", "")).strip()

        if not username:
            self.send_json(400, {"error": "Username cannot be empty."})
            return
        if not password:
            self.send_json(400, {"error": "Password cannot be empty."})
            return
        if len(username) > 15:
            self.send_json(400, {"error": "Username must be 15 characters or fewer."})
            return
        if len(password) > 15:
            self.send_json(400, {"error": "Password must be 15 characters or fewer."})
            return

        is_valid, error_message = self.validate_login(username, password)
        if not is_valid:
            self.register_failed_login(client_ip)
            self.send_json(401, {"error": error_message})
            return

        self.clear_failed_logins(client_ip)
        session_id = self.create_session(username)
        self.send_json(
            200,
            {"ok": True, "currentUser": username},
            {"Set-Cookie": self.build_session_cookie(session_id)},
        )

    def handle_logout(self):
        """Clear server session and cookie."""
        self.destroy_session()
        self.send_json(200, {"ok": True}, {"Set-Cookie": self.build_clear_cookie()})

    def handle_session_status(self):
        """Return current auth state."""
        session = self.get_authenticated_session()
        if not session:
            self.send_json(200, {"authenticated": False})
            return
        self.send_json(200, {"authenticated": True, "currentUser": session["username"]})

    def handle_config_request(self, route):
        """Serve JSON config files through authenticated APIs."""
        file_path = self.CONFIG_ENDPOINTS.get(route)
        if not file_path:
            self.send_json(404, {"error": "Config endpoint not found"})
            return

        if not file_path.exists():
            self.send_json(500, {"error": f"Missing config file: {file_path.name}"})
            return

        try:
            payload = json.loads(file_path.read_text(encoding="utf-8"))
            self.send_json(200, payload)
        except Exception:
            self.send_json(500, {"error": f"Failed to read {file_path.name}"})

    def handle_agenda_request(self):
        """Serve agenda extracted from Excel."""
        try:
            data = self.get_agenda_from_excel()
            self.send_json(200, data)
        except FileNotFoundError as err:
            self.send_json(500, {"error": str(err)})
        except Exception:
            self.send_json(500, {"error": "Unable to read agenda data"})

    def require_auth(self):
        """Enforce authenticated session for protected APIs."""
        if not self.get_authenticated_session():
            self.send_json(401, {"error": "Unauthorized"})
            return False
        return True

    def get_agenda_from_excel(self):
        """Read Excel file and return JSON data."""
        agenda_path = resolve_latest_agenda_path()
        wb = openpyxl.load_workbook(agenda_path, data_only=True)
        try:
            # Read column config from Config sheet (case-insensitive match)
            column_config = {}
            config_sheet_name = next((name for name in wb.sheetnames if normalize_label(name) == "config"), None)
            if config_sheet_name:
                config_ws = wb[config_sheet_name]
                for row_idx in range(2, config_ws.max_row + 1):
                    sheet_name = config_ws.cell(row_idx, 1).value
                    columns = config_ws.cell(row_idx, 2).value
                    if sheet_name and columns:
                        column_config[normalize_label(sheet_name)] = [
                            normalize_label(c) for c in str(columns).split(",") if normalize_label(c)
                        ]
            configured_sheets = set(column_config.keys())

            all_sheets_data = {}

            for sheet_name in wb.sheetnames:
                # Skip Config sheet
                if normalize_label(sheet_name) == "config":
                    continue
                # If config sheet has entries, only include configured sheets.
                if configured_sheets and normalize_label(sheet_name) not in configured_sheets:
                    continue

                ws = wb[sheet_name]
                if ws.sheet_state == "hidden":
                    continue

                header_row = detect_header_row(ws)

                # Get all headers
                all_headers = []
                header_map = {}
                for col_idx in range(1, ws.max_column + 1):
                    cell = ws.cell(row=header_row, column=col_idx)
                    header = str(cell.value).strip() if cell.value else ""
                    all_headers.append(header)
                    if header:
                        header_map[normalize_label(header)] = {"name": header, "index": col_idx}

                # Filter headers based on config
                sheet_key = normalize_label(sheet_name)
                if sheet_key in column_config:
                    headers = [
                        header_map[h]["name"]
                        for h in column_config[sheet_key]
                        if h in header_map
                    ]
                else:
                    headers = [h for h in all_headers if h]

                # Get data rows
                rows = []
                for row_idx in range(header_row + 1, ws.max_row + 1):
                    row_data = {}
                    for header in headers:
                        mapped = header_map.get(normalize_label(header))
                        if mapped:
                            cell = ws.cell(row=row_idx, column=mapped["index"])
                            row_data[header] = self.serialize_value(cell.value)
                    # Skip rows that are fully empty across configured headers.
                    if any(value != "" for value in row_data.values()):
                        rows.append(row_data)

                all_sheets_data[sheet_name] = {
                    "sheet_name": sheet_name,
                    "headers": headers,
                    "total_rows": len(rows),
                    "rows": rows,
                }

            return all_sheets_data
        finally:
            wb.close()

    def serialize_value(self, value):
        """Convert Excel values to JSON-serializable format."""
        if value is None:
            return ""
        if isinstance(value, datetime):
            if value.hour == 0 and value.minute == 0 and value.second == 0:
                return value.strftime("%Y-%m-%d")
            return value.strftime("%Y-%m-%d %H:%M")
        if isinstance(value, (int, float)):
            return value
        return str(value).strip()

    def read_json_body(self):
        """Read and parse JSON request body."""
        content_length = self.headers.get("Content-Length")
        if not content_length:
            return {}

        try:
            body_size = int(content_length)
        except ValueError as err:
            raise ValueError("Invalid Content-Length header") from err

        if body_size <= 0:
            return {}

        raw_body = self.rfile.read(body_size)
        if not raw_body:
            return {}

        try:
            return json.loads(raw_body.decode("utf-8"))
        except json.JSONDecodeError as err:
            raise ValueError("Invalid JSON payload") from err

    def send_json(self, status_code, payload, extra_headers=None):
        """Send JSON response."""
        body = json.dumps(payload).encode("utf-8")
        self.send_response(status_code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate")
        self.send_header("Content-Length", str(len(body)))
        if extra_headers:
            for header_name, header_value in extra_headers.items():
                self.send_header(header_name, header_value)
        try:
            self.end_headers()
            self.wfile.write(body)
        except (BrokenPipeError, ConnectionResetError, OSError):
            # Client disconnected before response write completed.
            return

    def get_client_ip(self):
        """Resolve client IP, including proxy-forwarded header."""
        forwarded = self.headers.get("X-Forwarded-For")
        if forwarded:
            return forwarded.split(",")[0].strip()
        return self.client_address[0]

    def get_credentials(self):
        """Load credentials from secure server-side file."""
        credentials_path = DATA_DIR / "credentials.json"
        if not credentials_path.exists():
            return []

        try:
            data = json.loads(credentials_path.read_text(encoding="utf-8"))
        except Exception:
            return []

        users = data.get("users", [])
        return users if isinstance(users, list) else []

    def validate_login(self, username, password):
        """Validate credentials while preserving existing user feedback."""
        users = self.get_credentials()

        for user in users:
            stored_username = str(user.get("username", ""))
            if not hmac.compare_digest(stored_username, username):
                continue

            if self.verify_password(password, user):
                return True, None
            return False, "Your password is invalid."

        return False, "Your username is invalid."

    def verify_password(self, provided_password, user):
        """Verify plain text or PBKDF2 hash password formats."""
        stored_hash = user.get("password_hash")
        if stored_hash:
            return self.verify_pbkdf2_hash(provided_password, str(stored_hash))

        stored_password = user.get("password")
        if stored_password is None:
            return False
        return hmac.compare_digest(str(stored_password), provided_password)

    def verify_pbkdf2_hash(self, provided_password, stored_hash):
        """
        Verify password hash format:
        pbkdf2_sha256$<iterations>$<salt>$<hex_digest>
        """
        parts = stored_hash.split("$")
        if len(parts) != 4 or parts[0] != "pbkdf2_sha256":
            return False

        try:
            iterations = int(parts[1])
        except ValueError:
            return False

        salt = parts[2]
        expected_digest = parts[3]
        computed_digest = hashlib.pbkdf2_hmac(
            "sha256",
            provided_password.encode("utf-8"),
            salt.encode("utf-8"),
            iterations,
        ).hex()
        return hmac.compare_digest(computed_digest, expected_digest)

    def is_rate_limited(self, client_ip):
        """Simple per-IP login rate limit."""
        now = time.time()
        with self.AUTH_LOCK:
            attempts = self.FAILED_LOGINS.get(client_ip, [])
            attempts = [ts for ts in attempts if now - ts < self.LOGIN_WINDOW_SECONDS]
            self.FAILED_LOGINS[client_ip] = attempts
            return len(attempts) >= self.LOGIN_MAX_ATTEMPTS

    def register_failed_login(self, client_ip):
        """Record failed login attempts for rate limit checks."""
        now = time.time()
        with self.AUTH_LOCK:
            attempts = self.FAILED_LOGINS.get(client_ip, [])
            attempts = [ts for ts in attempts if now - ts < self.LOGIN_WINDOW_SECONDS]
            attempts.append(now)
            self.FAILED_LOGINS[client_ip] = attempts

    def clear_failed_logins(self, client_ip):
        """Clear failed login state after successful login."""
        with self.AUTH_LOCK:
            self.FAILED_LOGINS.pop(client_ip, None)

    def create_session(self, username):
        """Create a new authenticated session."""
        session_id = secrets.token_urlsafe(48)
        expires_at = datetime.now(timezone.utc) + timedelta(minutes=self.SESSION_TTL_MINUTES)

        with self.AUTH_LOCK:
            self.cleanup_expired_sessions()
            self.SESSIONS[session_id] = {
                "username": username,
                "expires_at": expires_at,
            }
        return session_id

    def get_authenticated_session(self):
        """Return active session data, refreshing session expiry."""
        session_id = self.get_session_id_from_cookie()
        if not session_id:
            return None

        with self.AUTH_LOCK:
            self.cleanup_expired_sessions()
            session = self.SESSIONS.get(session_id)
            if not session:
                return None

            session["expires_at"] = datetime.now(timezone.utc) + timedelta(minutes=self.SESSION_TTL_MINUTES)
            return {"session_id": session_id, "username": session["username"]}

    def destroy_session(self):
        """Remove active session if present."""
        session_id = self.get_session_id_from_cookie()
        if not session_id:
            return
        with self.AUTH_LOCK:
            self.SESSIONS.pop(session_id, None)

    def cleanup_expired_sessions(self):
        """Purge expired sessions."""
        now = datetime.now(timezone.utc)
        expired_session_ids = [sid for sid, data in self.SESSIONS.items() if data["expires_at"] <= now]
        for session_id in expired_session_ids:
            self.SESSIONS.pop(session_id, None)

    def get_session_id_from_cookie(self):
        """Read session cookie from request headers."""
        cookie_header = self.headers.get("Cookie")
        if not cookie_header:
            return None

        cookie = SimpleCookie()
        cookie.load(cookie_header)
        morsel = cookie.get(self.SESSION_COOKIE_NAME)
        return morsel.value if morsel else None

    def build_session_cookie(self, session_id):
        """Build secure session cookie string."""
        cookie = SimpleCookie()
        cookie[self.SESSION_COOKIE_NAME] = session_id
        morsel = cookie[self.SESSION_COOKIE_NAME]
        morsel["path"] = "/"
        morsel["httponly"] = True
        morsel["samesite"] = "Lax"
        morsel["max-age"] = str(self.SESSION_TTL_MINUTES * 60)
        if self.COOKIE_SECURE:
            morsel["secure"] = True
        return morsel.OutputString()

    def build_clear_cookie(self):
        """Build cookie string that removes session cookie."""
        cookie = SimpleCookie()
        cookie[self.SESSION_COOKIE_NAME] = ""
        morsel = cookie[self.SESSION_COOKIE_NAME]
        morsel["path"] = "/"
        morsel["httponly"] = True
        morsel["samesite"] = "Lax"
        morsel["max-age"] = "0"
        morsel["expires"] = "Thu, 01 Jan 1970 00:00:00 GMT"
        if self.COOKIE_SECURE:
            morsel["secure"] = True
        return morsel.OutputString()

    def log_message(self, format, *args):
        """Reduce noisy static logs while keeping page/API hits visible."""
        if self.path.startswith("/assets/"):
            return
        if self.path.endswith(".json"):
            return
        if self.path.startswith("/api/"):
            print(f"API {self.command} {self.path}")
            return
        print(self.path)


if __name__ == "__main__":
    host = os.getenv("HOST", "0.0.0.0")
    port = read_env_int("PORT", 8000)

    server_address = (host, port)
    httpd = ThreadingHTTPServer(server_address, DynamicExcelHandler)

    print("Starting Travel Guide server...")
    print(f"Open: http://localhost:{port}")
    print("Session auth enabled; API data is protected")
    print("Press Ctrl+C to stop\n")

    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\nServer stopped")
